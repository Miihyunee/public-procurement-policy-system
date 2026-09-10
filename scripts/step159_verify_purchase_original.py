"""
STEP 159 — 실제 고객 구매실적 원본을 **임시 DB** 에 적재해 보는 검증 도구.

⛔ 원본을 열어 저장하지 않습니다(openpyxl read_only).
⛔ 고객 DB 를 쓰지 않습니다. 새 임시 DB 를 만듭니다.
⛔ 사업자등록번호·기업명을 그대로 찍지 않습니다(마스킹).
⛔ 저장소에 넣지 않습니다 — 임시 폴더에서만 돕니다.

사용법:
    python step159_verify.py "<원본 xlsx 경로>"
"""

from __future__ import annotations

import hashlib
import sqlite3
import sys
import tempfile
from collections import Counter
from datetime import date
from decimal import Decimal
from pathlib import Path

from fastapi.testclient import TestClient
from openpyxl import load_workbook

from procurement.app import create_app
from procurement.database.bootstrap import bootstrap


def rule(title: str) -> None:
    print()
    print("=" * 66)
    print(title)
    print("=" * 66)


def mask_name(value: object) -> str:
    text = str(value or "").strip()
    if len(text) <= 1:
        return "*"
    return text[0] + "*" * (len(text) - 1)


def mask_bizno(value: object) -> str:
    digits = "".join(ch for ch in str(value or "") if ch.isdigit())
    if len(digits) < 4:
        return "*" * len(digits)
    return digits[:3] + "*" * (len(digits) - 3)


def main(source: str) -> int:
    path = Path(source)

    # ── 1. 원본 확인 (읽기 전용) ────────────────────────────────────
    rule("1. 원본 파일")
    if not path.exists():
        print("파일이 없습니다:", path)
        return 1
    data = path.read_bytes()
    before_hash = hashlib.sha256(data).hexdigest()
    print("파일명 :", path.name)
    print("크기   :", f"{len(data):,} bytes")
    print("SHA256 :", before_hash)

    book = load_workbook(path, read_only=True, data_only=True)
    sheet = book.worksheets[0]
    print("시트   :", [ws.title for ws in book.worksheets])
    print("행/열  :", sheet.max_row, "/", sheet.max_column)
    rows = list(sheet.iter_rows(values_only=True))
    book.close()
    header = [str(v).strip() if v is not None else "" for v in rows[0]]
    body = [row for row in rows[1:] if any(cell is not None and str(cell).strip() for cell in row)]
    print("머리글 :", header)
    print("자료 행:", f"{len(body):,}")

    # ── 2. 원본만 보고 미리 세어 보기 ────────────────────────────────
    rule("2. 원본만 보고 미리 센 값 (프로그램을 거치기 전)")
    index = {name: position for position, name in enumerate(header)}
    for needed in ("계", "결의일자", "신고기준일", "거래처명", "사업자번호"):
        print(f"  {needed:<8}", "있음" if needed in index else "없음")

    amounts: list[Decimal] = []
    unreadable = 0
    for row in body:
        raw = row[index["계"]] if "계" in index else None
        # ⛔ ``raw or ""`` 로 쓰지 않는다 — 0원이 빈칸으로 둔갑한다.
        text = ("" if raw is None else str(raw)).replace(",", "").replace("₩", "").strip()
        try:
            amounts.append(Decimal(text))
        except Exception:  # noqa: BLE001 — 읽을 수 없는 값도 세기만 한다
            unreadable += 1
    negative = sum(1 for value in amounts if value < 0)
    zero = sum(1 for value in amounts if value == 0)
    positive = [value for value in amounts if value > 0]
    print(f"  금액 읽음 {len(amounts):,} · 읽지 못함 {unreadable:,}")
    print(f"  음수 {negative:,} · 0원 {zero:,} · 양수 {len(positive):,}")
    print(f"  양수 합계 {sum(positive):,}")

    years: Counter[str] = Counter()
    for row in body:
        raw = row[index["결의일자"]] if "결의일자" in index else None
        if isinstance(raw, date):
            years[str(raw.year)] += 1
        else:
            text = str(raw or "").strip()
            years[text[:4] if len(text) >= 4 else "(읽기 어려움)"] += 1
    print("  결의일자 연도별:", dict(sorted(years.items())))

    if len(years) != 1:
        print("  ⚠️ 연도가 하나가 아닙니다. 업로드 기간을 어떻게 잡을지 판단이 필요합니다.")
    year = int(next(iter(years)))

    # ── 3. 임시 DB ─────────────────────────────────────────────────
    rule("3. 임시 DB")
    workdir = Path(tempfile.mkdtemp(prefix="step159-"))
    db = workdir / "step159_purchase_verification.db"
    bootstrap(db)
    print("임시 DB :", db)
    print("⛔ 고객 DB 를 쓰지 않았습니다.")
    client = TestClient(create_app(db))

    # ── 4. 검증 ────────────────────────────────────────────────────
    rule("4. 검증 (저장하지 않음)")
    response = client.post("/uploads/purchases/validate", json={"file_path": str(path)})
    print("HTTP", response.status_code)
    result = response.json()
    for key in ("ok", "sheet_name", "total_rows", "valid_rows", "error_rows", "file_errors"):
        print(f"  {key:<12}", result.get(key))
    issues = result.get("issues") or []
    print(f"  issues       {len(issues):,}건")
    kinds: Counter[str] = Counter()
    for issue in issues:
        if isinstance(issue, dict):
            label = f"{issue.get('severity', '?')} | {issue.get('header')} | "
            label += str(issue.get("message", ""))[:44]
        else:
            label = str(issue)[:60]
        kinds[label] += 1
    for text, count in kinds.most_common(8):
        print(f"    {count:>6,}  {text}")

    if not result.get("ok"):
        print("⚠️ 검증을 통과하지 못했습니다. 적재로 넘어가지 않습니다.")
        return 2

    # ── 5. 적재 ────────────────────────────────────────────────────
    rule(f"5. 임시 DB 적재 (대상 연도 {year}, 월 지정 없음)")
    response = client.post(
        "/uploads/purchases", json={"file_path": str(path), "year": year}
    )
    print("HTTP", response.status_code)
    stored = response.json()
    for key in (
        "ok", "stored", "total_rows", "valid_rows", "error_rows",
        "stored_rows", "rejected_rows", "unexplained_rows", "batch_id",
    ):
        print(f"  {key:<16}", stored.get(key))
    for reason in stored.get("rejection_reasons") or []:
        print("   미적재:", reason)

    # ── 6. 저장된 내용 ──────────────────────────────────────────────
    rule("6. 저장된 내용 (임시 DB 직접 조회)")
    con = sqlite3.connect(db)
    count, total = con.execute("SELECT COUNT(*), COALESCE(SUM(amount),0) FROM purchase").fetchone()
    print(f"  purchase        {count:,}건")
    print(f"  amount 합계     {total:,}원")
    for column in ("contract_date", "payment_date"):
        filled = con.execute(
            f"SELECT COUNT(*) FROM purchase WHERE {column} IS NOT NULL"  # noqa: S608
        ).fetchone()[0]
        print(f"  {column:<15} 값이 있는 행 {filled:,} (기대: 0)")
    for column in ("company_name", "business_no", "resolution_date", "issue_date", "amount"):
        blank = con.execute(
            f"SELECT COUNT(*) FROM purchase WHERE {column} IS NULL OR TRIM({column})=''"  # noqa: S608
        ).fetchone()[0]
        print(f"  {column:<15} 비어 있는 행 {blank:,}")
    lo, hi = con.execute(
        "SELECT MIN(resolution_date), MAX(resolution_date) FROM purchase"
    ).fetchone()
    print("  결의일자 범위   ", lo, "~", hi)
    lo, hi = con.execute("SELECT MIN(issue_date), MAX(issue_date) FROM purchase").fetchone()
    print("  신고기준일 범위 ", lo, "~", hi)

    print("\n  표본 3건 (⛔ 마스킹):")
    for row in con.execute(
        "SELECT company_name, business_no, resolution_date, issue_date, amount, "
        "contract_date, payment_date, budget_account FROM purchase LIMIT 3"
    ):
        print(
            f"    {mask_name(row[0]):<10} {mask_bizno(row[1]):<12} 결의 {row[2]} "
            f"신고 {row[3]} 금액 {row[4]:,} 계약 {row[5]} 지급 {row[6]} 과목 {row[7]}"
        )

    rule("7. 배치")
    for row in con.execute(
        "SELECT batch_id, status, period_start, period_end FROM import_batch ORDER BY batch_id"
    ):
        print("  ", row)

    rule("8. 대시보드 API")
    status = client.get("/dashboard/data-status").json()
    for key in (
        "purchase_count", "purchase_total_amount", "matched_purchase_count",
        "unmatched_purchase_count", "batch_count", "active_batch_count",
        "superseded_batch_count", "calculation_target_count",
    ):
        print(f"  {key:<26}", status.get(key))

    rule("9. 원본이 그대로인가")
    after_hash = hashlib.sha256(path.read_bytes()).hexdigest()
    print("  전 :", before_hash)
    print("  후 :", after_hash)
    print("  동일:", before_hash == after_hash)

    print("\n임시 DB 는 여기 있습니다(끝나면 지우셔도 됩니다):")
    print(" ", workdir)
    return 0


if __name__ == "__main__":
    if len(sys.argv) != 2:
        print(__doc__)
        raise SystemExit(1)
    raise SystemExit(main(sys.argv[1]))
