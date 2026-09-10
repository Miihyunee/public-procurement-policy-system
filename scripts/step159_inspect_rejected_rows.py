"""
STEP 159 (후속) — 검증에서 걸린 13행이 무엇인지, 그리고 정상 2,292행만의
합계가 얼마인지 **읽기만** 해서 확인합니다.

⛔ 원본을 열어 저장하지 않습니다(openpyxl read_only).
⛔ DB 를 만들지도 건드리지도 않습니다.
⛔ 기업명·사업자등록번호를 그대로 찍지 않습니다(마스킹).
⛔ 원본을 고치지 않습니다 — 무엇이 들어 있는지 보기만 합니다.

사용법:
    python step159_inspect_rejected_rows.py "<원본 xlsx 경로>"
"""

from __future__ import annotations

import hashlib
import sys
from collections import Counter
from datetime import date, datetime
from decimal import Decimal
from pathlib import Path

from openpyxl import load_workbook

#: 이 넷이 모두 비어 있으면 프로그램이 그 행을 오류로 봅니다.
KEY_COLUMNS = ("결의일자", "신고기준일", "거래처명", "사업자번호")


def rule(title: str) -> None:
    print()
    print("=" * 66)
    print(title)
    print("=" * 66)


def mask_name(value: object) -> str:
    text = str(value if value is not None else "").strip()
    if not text:
        return "(빈칸)"
    if len(text) == 1:
        return "*"
    return text[0] + "*" * (len(text) - 1)


def mask_bizno(value: object) -> str:
    digits = "".join(ch for ch in str(value if value is not None else "") if ch.isdigit())
    if not digits:
        return "(빈칸)"
    if len(digits) < 4:
        return "*" * len(digits)
    return digits[:3] + "*" * (len(digits) - 3)


def as_amount(raw: object) -> Decimal | None:
    if raw is None:
        return None
    text = str(raw).replace(",", "").replace("₩", "").strip()
    if not text:
        return None
    try:
        return Decimal(text)
    except Exception:  # noqa: BLE001 — 읽히지 않는 값은 없음으로 센다
        return None


def as_year(raw: object) -> str:
    if isinstance(raw, date | datetime):
        return str(raw.year)
    text = str(raw if raw is not None else "").strip()
    return text[:4] if len(text) >= 4 else "(빈칸)"


def blank(raw: object) -> bool:
    return raw is None or not str(raw).strip()


def main(source: str) -> int:
    path = Path(source)
    if not path.exists():
        print("파일이 없습니다:", path)
        return 1

    before = hashlib.sha256(path.read_bytes()).hexdigest()
    book = load_workbook(path, read_only=True, data_only=True)
    sheet = book.worksheets[0]
    rows = list(sheet.iter_rows(values_only=True))
    book.close()

    header = [str(v).strip() if v is not None else "" for v in rows[0]]
    at = {name: position for position, name in enumerate(header)}
    body = [
        (number, row)
        for number, row in enumerate(rows[1:], start=2)
        if any(cell is not None and str(cell).strip() for cell in row)
    ]

    def cell(row: tuple[object, ...], name: str) -> object:
        position = at.get(name)
        return None if position is None else row[position]

    broken = [(n, r) for n, r in body if all(blank(cell(r, c)) for c in KEY_COLUMNS)]
    good = [(n, r) for n, r in body if not all(blank(cell(r, c)) for c in KEY_COLUMNS)]

    rule("1. 파일")
    print("파일명 :", path.name)
    print("SHA256 :", before)
    print("자료 행:", f"{len(body):,}")

    rule(f"2. 프로그램이 오류로 본 행 — {len(broken)}건 (⛔ 마스킹)")
    print("  핵심 네 칸(결의일자·신고기준일·거래처명·사업자번호)이 모두 빈 행입니다.")
    print()
    print(f"  {'엑셀행':>6} {'번호':>8} {'적요':<24} {'계':>18} {'예산과목':<12}")
    print("  " + "-" * 74)
    broken_sum = Decimal(0)
    for number, row in broken:
        amount = as_amount(cell(row, "계"))
        broken_sum += amount or Decimal(0)
        note = str(cell(row, "적요") or "").strip()[:22]
        print(
            f"  {number:>6} {str(cell(row, '번호') or ''):>8} {note:<24} "
            f"{(f'{amount:,}' if amount is not None else '(빈칸)'):>18} "
            f"{str(cell(row, '예산과목') or '')[:12]:<12}"
        )
    print("  " + "-" * 74)
    print(f"  이 {len(broken)}행의 「계」 합계: {broken_sum:,}")

    rule("3. 나머지 정상 행만 센 값")
    amounts = [as_amount(cell(r, "계")) for _, r in good]
    negative = [a for a in amounts if a is not None and a < 0]
    zero = [a for a in amounts if a is not None and a == 0]
    positive = [a for a in amounts if a is not None and a > 0]
    print(f"  정상 행           {len(good):,}")
    print(f"  음수              {len(negative):,}")
    print(f"  0원               {len(zero):,}")
    print(f"  양수(저장 대상)    {len(positive):,}")
    print(f"  양수 합계          {sum(positive):,}")
    print()
    print("  기존 기준과 대조")
    print(f"    2,292건 ↔ {len(good):,}건")
    print(f"    2,162건 ↔ {len(positive):,}건")
    print(
        f"    130건   ↔ {len(negative) + len(zero):,}건 "
        f"(음수 {len(negative):,} · 0원 {len(zero):,})"
    )
    print(f"    9,582,813,023원 ↔ {sum(positive):,}원")

    rule("4. 결의일자 연도 (정상 행만)")
    years: Counter[str] = Counter(as_year(cell(r, "결의일자")) for _, r in good)
    for label, count in sorted(years.items()):
        print(f"  {label:<12} {count:>6,}")

    other = [(n, r) for n, r in good if as_year(cell(r, "결의일자")) != "2026"]
    if other:
        rule(f"5. 2026년이 아닌 결의일자 — {len(other)}건 (⛔ 마스킹)")
        for number, row in other:
            print(
                f"  {number:>6}행  결의 {cell(row, '결의일자')}  신고 {cell(row, '신고기준일')}  "
                f"거래처 {mask_name(cell(row, '거래처명'))}  "
                f"번호 {mask_bizno(cell(row, '사업자번호'))}  "
                f"계 {as_amount(cell(row, '계'))}"
            )

    rule("6. 원본이 그대로인가")
    after = hashlib.sha256(path.read_bytes()).hexdigest()
    print("  동일:", before == after)
    return 0


if __name__ == "__main__":
    if len(sys.argv) != 2:
        print(__doc__)
        raise SystemExit(1)
    raise SystemExit(main(sys.argv[1]))
